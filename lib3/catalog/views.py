from django.shortcuts import render, get_object_or_404
from django.db.models import Q
from rest_framework import viewsets, filters
from rest_framework.decorators import action
from rest_framework.response import Response

from .models import Book, BookCopy, Branch, Author, Genre
from .serializers import BookSerializer, BookListSerializer, BranchSerializer, AuthorSerializer


# ─── Web Views ────────────────────────────────────────────────

def book_list(request):
    query     = request.GET.get('q', '').strip()
    genre_id  = request.GET.get('genre', '')
    branch_id = request.GET.get('branch', '')

    books = Book.objects.prefetch_related('authors', 'copies').select_related('genre')

    if query:
        # Ищем по всем полям: название, имя автора, фамилия автора,
        # ISBN, жанр, описание, год
        books = books.filter(
            Q(title__icontains=query)              |
            Q(authors__first_name__icontains=query)|
            Q(authors__last_name__icontains=query) |
            Q(isbn__icontains=query)               |
            Q(genre__name__icontains=query)        |
            Q(description__icontains=query)        |
            Q(year__icontains=query)
        ).distinct()

    if genre_id:
        books = books.filter(genre_id=genre_id)

    if branch_id:
        books = books.filter(copies__branch_id=branch_id).distinct()

    context = {
        'books':           books,
        'genres':          Genre.objects.all(),
        'branches':        Branch.objects.filter(is_active=True),
        'query':           query,
        'selected_genre':  genre_id,
        'selected_branch': branch_id,
    }
    return render(request, 'catalog/book_list.html', context)


def book_detail(request, pk):
    book = get_object_or_404(
        Book.objects.prefetch_related('authors', 'copies__branch').select_related('genre'),
        pk=pk
    )
    return render(request, 'catalog/book_detail.html', {'book': book})


def add_copy(request, book_pk):
    book     = get_object_or_404(Book, pk=book_pk)
    branches = Branch.objects.filter(is_active=True)
    error    = None

    if request.method == 'POST':
        branch_id  = request.POST.get('branch_id')
        inv_number = request.POST.get('inventory_number', '').strip()

        if not branch_id or not inv_number:
            error = 'Լрацрек болор дашtерэ'
        elif BookCopy.objects.filter(inventory_number=inv_number).exists():
            error = f'Гуjqаин № «{inv_number}» арден гоjутюн уни'
        else:
            branch = get_object_or_404(Branch, pk=branch_id)
            BookCopy.objects.create(
                book=book, branch=branch,
                inventory_number=inv_number, status='available',
            )
            from django.shortcuts import redirect
            from django.contrib import messages
            messages.success(request, f'Оринак «{inv_number}» аvelацваd е')
            return redirect('book_detail', pk=book_pk)

    return render(request, 'catalog/add_copy.html', {
        'book': book, 'branches': branches, 'error': error,
    })


def branch_manage(request):
    from django.conf import settings
    branches      = Branch.objects.all()
    configured_dbs = set(settings.DATABASES.keys())
    return render(request, 'catalog/branch_manage.html', {
        'branches': branches, 'configured_dbs': configured_dbs,
    })


def branch_add(request):
    from django.contrib import messages as msg_fw
    if request.method == 'POST':
        name     = request.POST.get('name', '').strip()
        db_alias = request.POST.get('db_alias', '').strip()
        address  = request.POST.get('address', '').strip()
        phone    = request.POST.get('phone', '').strip()

        if not name or not db_alias:
            msg_fw.error(request, 'Анванумэ эв БД alias-э патрастатрваd eн')
        elif Branch.objects.filter(db_alias=db_alias).exists():
            msg_fw.error(request, f'«{db_alias}» alias-ов масначял арден ка')
        else:
            Branch.objects.create(
                name=name, db_alias=db_alias,
                address=address, phone=phone, is_active=True,
            )
            msg_fw.success(request, f'Масначял «{name}» ствараgопкваd е')

    from django.shortcuts import redirect
    return redirect('branch_manage')


# ─── API ViewSets ─────────────────────────────────────────────

class BookViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = Book.objects.prefetch_related('authors', 'copies').select_related('genre')
    filter_backends  = [filters.SearchFilter, filters.OrderingFilter]
    search_fields    = ['title', 'authors__last_name', 'authors__first_name', 'isbn', 'genre__name']
    ordering_fields  = ['title', 'year']

    def get_serializer_class(self):
        return BookListSerializer if self.action == 'list' else BookSerializer

    @action(detail=False, methods=['get'])
    def available(self, request):
        branch_id = request.query_params.get('branch_id')
        copies    = BookCopy.objects.filter(status='available')
        if branch_id:
            copies = copies.filter(branch_id=branch_id)
        book_ids  = copies.values_list('book_id', flat=True).distinct()
        books     = self.get_queryset().filter(id__in=book_ids)
        return Response(BookListSerializer(books, many=True).data)


class BranchViewSet(viewsets.ReadOnlyModelViewSet):
    queryset         = Branch.objects.filter(is_active=True)
    serializer_class = BranchSerializer


class AuthorViewSet(viewsets.ReadOnlyModelViewSet):
    queryset         = Author.objects.all()
    serializer_class = AuthorSerializer
    filter_backends  = [filters.SearchFilter]
    search_fields    = ['first_name', 'last_name']


# ─── Массовый импорт ──────────────────────────────────────────

import csv
import io
from django.contrib import messages
from django.shortcuts import redirect

def bulk_add_books(request):
    branches = Branch.objects.filter(is_active=True)
    genres   = Genre.objects.all()
    results  = None

    if request.method == 'POST':
        act = request.POST.get('action')
        if act == 'manual':
            results = _process_manual(request, branches)
        elif act == 'file':
            results = _process_file(request, branches)

    return render(request, 'catalog/bulk_add.html', {
        'branches': branches, 'genres': genres, 'results': results,
    })


def _get_or_create_authors(names_str):
    authors = []
    for name in names_str.split(','):
        name = name.strip()
        if not name:
            continue
        parts = name.rsplit(' ', 1)
        first, last = (parts[0], parts[1]) if len(parts) == 2 else ('', parts[0])
        author, _ = Author.objects.get_or_create(last_name=last, first_name=first)
        authors.append(author)
    return authors


def _save_book_row(row, default_branch):
    title = row.get('title', '').strip()
    if not title:
        return False, 'Бац тогнваd: вернагир чка'

    inv = row.get('inventory_number', '').strip()
    if inv in ('—', '-', '–'):
        inv = ''
    if inv and BookCopy.objects.filter(inventory_number=inv).exists():
        return False, f'«{title}» — гуjqаин № «{inv}» арден гоjутюн уни'

    genre      = None
    genre_name = row.get('genre', '').strip()
    if genre_name:
        genre, _ = Genre.objects.get_or_create(name=genre_name)

    isbn = row.get('isbn', '').strip()
    if isbn in ('—', '-', '–', ''):
        isbn = None

    book = None
    if isbn:
        book = Book.objects.filter(isbn=isbn).first()
    if not book:
        book = Book.objects.filter(title=title).first()
    if not book:
        book = Book.objects.create(
            title=title, isbn=isbn,
            year=row.get('year') or None, genre=genre,
        )
    elif genre:
        book.genre = genre
        book.save()

    authors_str = row.get('authors', '').strip()
    if authors_str:
        book.authors.set(_get_or_create_authors(authors_str))

    if inv:
        branch_id = row.get('branch_id') or (default_branch.pk if default_branch else None)
        branch    = Branch.objects.filter(pk=branch_id).first() or default_branch
        if branch:
            BookCopy.objects.create(
                book=book, branch=branch,
                inventory_number=inv, status='available'
            )
            return True, f'✅ «{title}» [#{inv}]'
        return True, f'✅ «{title}» (без экземпляра — филиал не выбран)'
    return True, f'✅ «{title}» (без экземпляра)'


def _process_manual(request, branches):
    titles            = request.POST.getlist('title')
    default_branch_id = request.POST.get('default_branch_id')
    default_branch    = Branch.objects.filter(pk=default_branch_id).first()
    results           = []

    for i, title in enumerate(titles):
        if not title.strip():
            continue
        def get(lst, idx): return lst[idx] if idx < len(lst) else ''
        row = {
            'title':            title,
            'authors':          get(request.POST.getlist('authors'), i),
            'genre':            get(request.POST.getlist('genre'), i),
            'year':             get(request.POST.getlist('year'), i),
            'isbn':             get(request.POST.getlist('isbn'), i),
            'inventory_number': get(request.POST.getlist('inventory_number'), i),
            'branch_id':        default_branch_id,
        }
        ok, msg = _save_book_row(row, default_branch)
        results.append({'ok': ok, 'msg': msg})
    return results


def _process_file(request, branches):
    f = request.FILES.get('csv_file')
    if not f:
        return [{'ok': False, 'msg': 'Файл чи ынтрваd'}]

    default_branch_id = request.POST.get('default_branch_id')
    default_branch    = Branch.objects.filter(pk=default_branch_id).first()
    name              = f.name.lower()
    results           = []

    if name.endswith('.xlsx'):
        try:
            import openpyxl
            wb      = openpyxl.load_workbook(f)
            ws      = wb.active
            headers = [str(c.value or '').strip().lower() for c in next(ws.iter_rows(max_row=1))]
            for row_cells in ws.iter_rows(min_row=2, values_only=True):
                if all(v is None for v in row_cells):
                    continue
                row = {headers[i]: str(v or '').strip() for i, v in enumerate(row_cells)}
                ok, msg = _save_book_row(row, default_branch)
                results.append({'ok': ok, 'msg': msg})
        except Exception as e:
            results.append({'ok': False, 'msg': f'Excel сxал: {e}'})

    elif name.endswith('.csv'):
        try:
            text   = f.read().decode('utf-8-sig')
            reader = csv.DictReader(io.StringIO(text))
            for row in reader:
                clean = {k.strip().lower(): v.strip() for k, v in row.items()}
                ok, msg = _save_book_row(clean, default_branch)
                results.append({'ok': ok, 'msg': msg})
        except Exception as e:
            results.append({'ok': False, 'msg': f'CSV сxал: {e}'})
    else:
        results.append({'ok': False, 'msg': 'Хармар eн мiаjн .csv кам .xlsx фаjлерэ'})

    return results


def csv_template(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from django.http import HttpResponse
    import io

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Гирkер'

    headers       = ['title', 'authors', 'genre', 'year', 'isbn', 'inventory_number']
    header_labels = ['Анванум *', 'Хегhинакнер', 'Жанр', 'Тари', 'ISBN', 'Гуjqаин №']
    header_fill   = PatternFill('solid', fgColor='2B4590')

    for col, (key, label) in enumerate(zip(headers, header_labels), 1):
        cell       = ws.cell(row=1, column=col, value=key)
        cell.font  = Font(bold=True, color='FFFFFF')
        cell.fill  = header_fill
        cell.alignment = Alignment(horizontal='center')
        hint       = ws.cell(row=2, column=col, value=label)
        hint.font  = Font(italic=True, color='888888')

    examples = [
        ['Варпетэ эв Маpгаритан', 'Булгаков Михаил', 'Веп',      '1967', '978-5-17-059208-0', 'INV-001'],
        ['1984',                   'Оруэл Джордж',    'Хакautopia','1949', '978-5-17-059209-7', 'INV-002'],
        ['Хари Потерэ',            'Роулинг Джоан',   'Фентези',  '1997', '978-5-389-07435-4', 'INV-003'],
    ]
    for row_i, row_data in enumerate(examples, 3):
        for col_i, val in enumerate(row_data, 1):
            ws.cell(row=row_i, column=col_i, value=val)

    for col, width in enumerate([35, 25, 15, 8, 22, 14], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    response = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="books_template.xlsx"'
    return response